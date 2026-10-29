"""
inspect_workbooks.py
====================
Inspects all original workbooks and raw data files in the EIS project workspace.
Preserves original files without modifying them.
"""

import os
import openpyxl
import pandas as pd

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

print("=" * 80)
print("INSPECTING WORKBOOKS AND RAW DATA")
print("=" * 80)

raw_files = [
    'E-WASTE VALUE.xlsx',
    'EWaste_Forecasting_Master_Data.xlsx',
    'RS_Session_266_AU_2384_A.csv'
]

for fname in raw_files:
    fpath = os.path.join(PROJECT_DIR, fname)
    if not os.path.exists(fpath):
        fpath = os.path.join(PROJECT_DIR, 'original_data', fname)
    
    print(f"\n>>> File: {fname} (Size: {os.path.getsize(fpath)} bytes)")
    if fname.endswith(('.xlsx', '.xls')):
        wb = openpyxl.load_workbook(fpath, data_only=True)
        print(f"    Sheets found ({len(wb.sheetnames)}): {wb.sheetnames}")
        for sname in wb.sheetnames:
            ws = wb[sname]
            print(f"      - Sheet '{sname}': {ws.max_row} rows, {ws.max_column} columns")
    elif fname.endswith('.csv'):
        df = pd.read_csv(fpath)
        print(f"    CSV Shape: {df.shape}")
        print(f"    Columns: {list(df.columns)}")
        print(df.to_string(index=False))

print("\nInspection complete. Original files remain untouched.")
